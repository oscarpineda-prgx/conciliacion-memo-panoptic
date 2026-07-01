from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from ..paths import PROJECT_ROOT
from .formatting import HEADER_ROWS, style_workbook, write_sheet


# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

_P1_EXCLUDE_PROJECT = "MEX_SORIANA_Scope_2019"
_P1_EXCLUDE_POSTING = "NS-EneAgo25"
_P1_EXCLUDE_CAUSE   = "Limitation in a systems functionality"
_P1_VALID_STAGES    = {"Posting", "Vendor"}
_P1_VALID_STATUSES  = {"In Review", "Posted"}
_AMOUNT_TOLERANCE   = 1.0

_P4_KEY_COLS = [
    "Num Proveedor", "Año", "Concepto", "Num Categoria",
    "Monto antes impuestos", "IEPS", "IVA",
]

_BULK_UPDATE_COLS = ["Project name", "Claim number", "Posting reference number"]

_INCON_MONTOS_COLS = [
    "Vendor number", "Vendor name", "Assigned to",
    "Sum Net claim amount (Panoptic)", "Sum Tax amount (Panoptic)", "Gross claim amount (Panoptic)",
    "Monto antes de impuestos (MEMO)", "IVA (MEMO)", "IEPS (MEMO)", "Monto Total (MEMO)",
    "Amount Difference", "Vendor Status", "Amount Status",
]
_DEFAULT_TEMPLATE_PATH = PROJECT_ROOT / "data" / "templates" / "Claim_Bulk_Update_Template.xlsx"


# ---------------------------------------------------------------------------
# Utilidades de carga
# ---------------------------------------------------------------------------

def _normalize_vendor(series: pd.Series) -> pd.Series:
    """Normaliza claves de proveedor: strip, quita decimales '4423.0' → '4423'."""
    return series.astype(str).str.strip().str.replace(r"\.0$", "", regex=True)


def _normalize_posting_ref(series: pd.Series) -> pd.Series:
    """Normaliza Posting reference number a formato canónico de 3 dígitos.

    Ejemplos: M45 → M045, M0045 → M045, M045 → M045.
    Valores que no siguen el patrón letra+número (ej. NS-EneAgo25) se dejan intactos.
    """
    _pattern = re.compile(r"^([A-Za-z]+)(\d+)$")

    def _norm(val: object) -> object:
        if val is None or (isinstance(val, float) and pd.isna(val)):
            return val
        s = str(val).strip()
        m = _pattern.match(s)
        if m:
            return f"{m.group(1).upper()}{int(m.group(2)):03d}"
        return s

    return series.map(_norm)


def _ensure_tax_column(df: pd.DataFrame) -> None:
    """Agrega Total tax amount si no existe (Claim amount − Net claim amount)."""
    if "Total tax amount" not in df.columns:
        if "Claim amount" in df.columns and "Net claim amount" in df.columns:
            df["Total tax amount"] = df["Claim amount"] - df["Net claim amount"]
        else:
            df["Total tax amount"] = 0.0


def _load_panoptic(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path)
    df["Vendor number"] = _normalize_vendor(df["Vendor number"])
    if "Posting reference number" in df.columns:
        df["Posting reference number"] = _normalize_posting_ref(df["Posting reference number"])
    if "Financial year of origin" in df.columns:
        df["Financial year of origin"] = pd.to_numeric(
            df["Financial year of origin"], errors="coerce"
        ).astype("Int64")
    return df


def _load_memo_prov_fase(memo_path: Path) -> pd.DataFrame:
    """Hoja MontoxProveedorxFase — encabezados reales en fila 7 (header=6)."""
    df = pd.read_excel(memo_path, sheet_name="MontoxProveedorxFase", header=6)
    cols_needed = [
        "Num Proveedor", "Proveedor", "Fase",
        "Monto antes impuestos", "IEPS", "IVA", "Monto Total",
    ]
    df = df[[c for c in cols_needed if c in df.columns]].copy()
    df = df.dropna(subset=["Num Proveedor"])
    df["Num Proveedor"] = _normalize_vendor(df["Num Proveedor"])
    for col in ["Monto antes impuestos", "IEPS", "IVA", "Monto Total"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    return df


def _load_memo_concepto_ano(memo_path: Path) -> pd.DataFrame:
    """Hoja MontoxConceptoxAño — encabezados reales en fila 2 (header=1).
    El nombre de la hoja puede tener espacios al final o variante de ñ."""
    xl = pd.ExcelFile(memo_path)
    sheet_name = next((s for s in xl.sheet_names if "Concepto" in s), None)
    if sheet_name is None:
        raise ValueError(f"No se encontró la hoja MontoxConceptoxAño en {memo_path}")

    df = pd.read_excel(memo_path, sheet_name=sheet_name, header=1)

    year_col = next(
        (c for c in df.columns if isinstance(c, str) and c.startswith("A") and "o" in c.lower()),
        None,
    )
    if year_col and year_col != "Año":
        df = df.rename(columns={year_col: "Año"})

    cols_needed = [
        "Num Proveedor", "Proveedor", "Año", "Negocio", "Fase",
        "Concepto", "Num Categoria", "Nombre Categoria",
        "Monto antes impuestos", "IEPS", "IVA",
    ]
    df = df[[c for c in cols_needed if c in df.columns]].copy()
    df = df.dropna(subset=["Num Proveedor"])
    df["Num Proveedor"] = _normalize_vendor(df["Num Proveedor"])
    df["Año"] = pd.to_numeric(df["Año"], errors="coerce").astype("Int64")
    for col in ["Monto antes impuestos", "IEPS", "IVA"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    return df


# ---------------------------------------------------------------------------
# Descubrimiento de archivos MEMO
# ---------------------------------------------------------------------------

def _memo_id_from_name(name: str) -> str | None:
    """Extrae y normaliza el ID de memo a formato canónico de 3 dígitos.

    Memo_40_xxx → M040, Memo_040_xxx → M040, Memo_0040_xxx → M040.
    """
    match = re.search(r"[Mm]emo_(\d+)", name, re.IGNORECASE)
    return f"M{int(match.group(1)):03d}" if match else None


def _memo_id_from_path(memo_path: Path) -> str | None:
    return _memo_id_from_name(memo_path.stem)


def _find_memo_xlsx(folder: Path) -> Path | None:
    """Devuelve el Excel a usar de una carpeta de memo.
    Prioridad: archivo con 'correc' en el nombre; si no, el más reciente.
    """
    candidates = [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")]
    if not candidates:
        return None
    correccion = [p for p in candidates if re.search(r"correc", p.stem, re.IGNORECASE)]
    if correccion:
        return max(correccion, key=lambda p: p.stat().st_mtime)
    return max(candidates, key=lambda p: p.stat().st_mtime)


# ---------------------------------------------------------------------------
# Paso 1 — Asignación de Posting reference a claims sin memo
# ---------------------------------------------------------------------------

def _p1_filter_mask(
    df: pd.DataFrame,
    extra_stages: set[str] | None = None,
    extra_statuses: set[str] | None = None,
) -> pd.Series:
    """Máscara de los 5 filtros P1: claims candidatos a ser asignados a un memo.

    extra_stages:   stages adicionales a incluir además de _P1_VALID_STAGES (ej. {"Invoice"}).
    extra_statuses: statuses adicionales a incluir además de _P1_VALID_STATUSES
                    (ej. {"Invoice ready", "Invoiced"}).
    """
    def _str(col: str) -> pd.Series:
        return df[col].fillna("").astype(str).str.strip()

    valid_stages   = _P1_VALID_STAGES   | (extra_stages   or set())
    valid_statuses = _P1_VALID_STATUSES | (extra_statuses or set())

    mask = pd.Series(True, index=df.index)
    if "Project name" in df.columns:
        mask &= _str("Project name") != _P1_EXCLUDE_PROJECT
    if "Posting reference number" in df.columns:
        mask &= _str("Posting reference number") != _P1_EXCLUDE_POSTING
    if "Claim cause description" in df.columns:
        mask &= _str("Claim cause description") != _P1_EXCLUDE_CAUSE
    if "Stage" in df.columns:
        mask &= df["Stage"].isin(valid_stages)
    if "Status" in df.columns:
        mask &= df["Status"].isin(valid_statuses)
    return mask


_P1_COLS = [
    "Vendor number", "Proveedor (MEMO)", "Año",
    "Net claim amount", "Total tax amount", "Gross Panoptic (P1)",
    "Monto antes impuestos", "IEPS", "IVA", "Monto Total (MEMO)",
    "Diferencia",
]


def _paso1_asignar_memo(
    df_pan: pd.DataFrame,
    df_memo_ca: pd.DataFrame,
    memo_id: str,
    memo_vendor_sets: dict[str, set[str]],
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """P1: Identifica claims sin Posting reference y asigna memo_id si el monto coincide.

    Cruza Panoptic vs MontoxConceptoxAño por proveedor + año.
    Modifica df_pan in-place actualizando 'Posting reference number'.
    Para claims que ya tienen un memo diferente asignado, los sobreescribe solo si ese
    vendor NO existe en el MEMO de ese otro memo (asignación incorrecta en Panoptic).
    Devuelve (df_actualizados, df_diferencias, df_claims).
    """
    p1_mask = _p1_filter_mask(df_pan)
    year_col = "Financial year of origin"

    p1_by_vendor_year = (
        df_pan[p1_mask]
        .groupby(["Vendor number", year_col])
        .agg(net=("Net claim amount", "sum"), tax=("Total tax amount", "sum"))
    )
    p1_by_vendor_year["Gross"] = p1_by_vendor_year["net"] + p1_by_vendor_year["tax"]

    memo_totals = (
        df_memo_ca
        .groupby(["Num Proveedor", "Año"])
        .agg(
            proveedor=("Proveedor", "first"),
            monto_antes=("Monto antes impuestos", "sum"),
            ieps=("IEPS", "sum"),
            iva=("IVA", "sum"),
        )
    )
    memo_totals["Monto Total"] = (
        memo_totals["monto_antes"] + memo_totals["ieps"] + memo_totals["iva"]
    )

    actualizados:   list[dict] = []
    diferencias:    list[dict] = []
    all_claim_rows: list[pd.DataFrame] = []

    for (vendor, ano), memo_row in memo_totals.iterrows():
        memo_total = float(memo_row["Monto Total"])
        base = {
            "Vendor number":         vendor,
            "Proveedor (MEMO)":      memo_row["proveedor"],
            "Año":                   ano,
            "Monto antes impuestos": float(memo_row["monto_antes"]),
            "IEPS":                  float(memo_row["ieps"]),
            "IVA":                   float(memo_row["iva"]),
            "Monto Total (MEMO)":    memo_total,
        }

        key = (vendor, ano)
        if key in p1_by_vendor_year.index:
            pan_row   = p1_by_vendor_year.loc[key]
            pan_net   = float(pan_row["net"])
            pan_tax   = float(pan_row["tax"])
            pan_gross = float(pan_row["Gross"])
            diff = memo_total - pan_gross
            row = {
                **base,
                "Net claim amount":    pan_net,
                "Total tax amount":    pan_tax,
                "Gross Panoptic (P1)": pan_gross,
                "Diferencia":          diff,
            }

            if abs(diff) <= _AMOUNT_TOLERANCE:
                update_mask = (
                    p1_mask
                    & (df_pan["Vendor number"] == vendor)
                    & (df_pan[year_col] == ano)
                )

                # Determinar qué filas dentro de update_mask deben asignarse:
                # 1. Filas con Posting reference vacío → siempre asignar.
                # 2. Filas con un memo diferente asignado → asignar solo si ese vendor
                #    NO existe en el MEMO de ese otro memo (era una asignación incorrecta).
                # 3. Filas que ya tienen memo_id correcto → no tocar.
                update_idx = df_pan.index[update_mask]
                posting_refs = (
                    df_pan.loc[update_idx, "Posting reference number"]
                    .fillna("").astype(str).str.strip()
                )

                def _should_assign(ref: str) -> bool:
                    if ref == "":
                        return True   # vacío → siempre asignar
                    if ref == memo_id:
                        return False  # ya correcto → no tocar
                    # Memo diferente: asignar solo si el vendor NO está en ese MEMO
                    vendor_set = memo_vendor_sets.get(ref)
                    if vendor_set is None:
                        return False  # memo desconocido → conservador, no tocar
                    return vendor not in vendor_set

                should_assign = posting_refs.map(_should_assign)
                assign_idx = update_idx[should_assign]
                assign_mask = pd.Series(False, index=df_pan.index)
                assign_mask.iloc[df_pan.index.get_indexer(assign_idx)] = True

                df_pan.loc[assign_idx, "Posting reference number"] = memo_id
                actualizados.append(row)

                claim_cols = [c for c in _BULK_UPDATE_COLS if c in df_pan.columns]
                if claim_cols and len(assign_idx) > 0:
                    affected = df_pan.loc[assign_idx, claim_cols].copy()
                    all_claim_rows.append(affected)
            else:
                diferencias.append(row)
        else:
            diferencias.append({
                **base,
                "Net claim amount":    0.0,
                "Total tax amount":    0.0,
                "Gross Panoptic (P1)": 0.0,
                "Diferencia":          memo_total,
            })

    df_claims = (
        pd.concat(all_claim_rows, ignore_index=True)
        if all_claim_rows
        else pd.DataFrame(columns=_BULK_UPDATE_COLS)
    )

    # Siempre usar columnas definidas para que write_sheet escriba el encabezado
    # aunque la lista esté vacía (evita hojas con menos de HEADER_ROWS filas).
    return (
        pd.DataFrame(actualizados, columns=_P1_COLS),
        pd.DataFrame(diferencias,  columns=_P1_COLS),
        df_claims,
    )


def _paso1_reejecutar_diferencias(
    df_pan: pd.DataFrame,
    df_memo_ca: pd.DataFrame,
    df_diferencias: pd.DataFrame,
) -> pd.DataFrame:
    """Segunda pasada de P1: re-cruza los pares (vendor, año) que quedaron en P1 Diferencias
    añadiendo Stage='Invoice' al filtro. Devuelve solo los que ahora cuadran (|diff| <= tolerancia).

    No modifica df_pan ni asigna Posting reference — es solo informativa.
    """
    if df_diferencias.empty:
        return pd.DataFrame(columns=[*_P1_COLS, "Estado"])

    year_col = "Financial year of origin"
    p1_mask_ext = _p1_filter_mask(
        df_pan,
        extra_stages={"Invoice"},
        extra_statuses={"Invoice ready", "Invoiced"},
    )

    p1_ext = (
        df_pan[p1_mask_ext]
        .groupby(["Vendor number", year_col])
        .agg(net=("Net claim amount", "sum"), tax=("Total tax amount", "sum"))
    )
    p1_ext["Gross"] = p1_ext["net"] + p1_ext["tax"]

    recuperados: list[dict] = []

    for _, dif_row in df_diferencias.iterrows():
        vendor     = dif_row["Vendor number"]
        ano        = dif_row["Año"]
        memo_total = float(dif_row["Monto Total (MEMO)"])

        key = (vendor, ano)
        if key not in p1_ext.index:
            continue

        pan_row   = p1_ext.loc[key]
        pan_net   = float(pan_row["net"])
        pan_tax   = float(pan_row["tax"])
        pan_gross = float(pan_row["Gross"])
        diff = memo_total - pan_gross

        if abs(diff) <= _AMOUNT_TOLERANCE:
            recuperados.append({
                "Vendor number":         vendor,
                "Proveedor (MEMO)":      dif_row["Proveedor (MEMO)"],
                "Año":                   ano,
                "Net claim amount":      pan_net,
                "Total tax amount":      pan_tax,
                "Gross Panoptic (P1)":   pan_gross,
                "Monto antes impuestos": dif_row["Monto antes impuestos"],
                "IEPS":                  dif_row["IEPS"],
                "IVA":                   dif_row["IVA"],
                "Monto Total (MEMO)":    memo_total,
                "Diferencia":            diff,
                "Estado":                "Cruce exitoso",
            })

    return pd.DataFrame(recuperados, columns=[*_P1_COLS, "Estado"])


# ---------------------------------------------------------------------------
# Paso 2 + 3 — Cruce resumen (inconsistencias de proveedor y monto)
# ---------------------------------------------------------------------------

def _cruce_resumen(df_pan_memo: pd.DataFrame, df_memo_pf: pd.DataFrame) -> pd.DataFrame:
    """Outer join a nivel proveedor entre Panoptic (filtrado por memo_id) y MEMO.

    Columnas de salida: Vendor Status (Match / Missing in MEMO / Missing in Panoptic)
    y Amount Status (OK / Mismatch).
    """
    # Excluir claims de "Limitation in a systems functionality" del cálculo de montos
    if "Claim cause description" in df_pan_memo.columns:
        df_pan_memo = df_pan_memo[
            df_pan_memo["Claim cause description"].fillna("").str.strip() != _P1_EXCLUDE_CAUSE
        ]

    agg_cols: dict = {"Net claim amount": "sum", "Total tax amount": "sum"}
    if "Assigned to" in df_pan_memo.columns:
        agg_cols["Assigned to"] = "first"

    pan_summary = (
        df_pan_memo
        .groupby(["Vendor number", "Vendor name"], as_index=False)
        .agg(agg_cols)
    )
    pan_summary["Gross claim amount (Panoptic)"] = (
        pan_summary["Net claim amount"] + pan_summary["Total tax amount"]
    )
    pan_summary = pan_summary.rename(columns={
        "Net claim amount": "Sum Net claim amount (Panoptic)",
        "Total tax amount": "Sum Tax amount (Panoptic)",
    })

    _memo_agg_cols = {
        col: "sum"
        for col in ["Monto antes impuestos", "IEPS", "IVA", "Monto Total"]
        if col in df_memo_pf.columns
    }
    memo_summary = (
        df_memo_pf
        .groupby(["Num Proveedor", "Proveedor"], as_index=False)
        .agg(_memo_agg_cols)
        .rename(columns={
            "Num Proveedor":        "Vendor number",
            "Proveedor":            "Vendor name (MEMO)",
            "Monto antes impuestos":"Monto antes de impuestos (MEMO)",
            "IEPS":                 "IEPS (MEMO)",
            "IVA":                  "IVA (MEMO)",
            "Monto Total":          "Monto Total (MEMO)",
        })
    )

    cruce = pd.merge(pan_summary, memo_summary, on="Vendor number", how="outer", indicator=True)
    cruce["Vendor Status"] = np.where(
        cruce["_merge"] == "both",        "Match",
        np.where(cruce["_merge"] == "left_only", "Missing in MEMO", "Missing in Panoptic"),
    )
    cruce["Gross claim amount (Panoptic)"] = cruce["Gross claim amount (Panoptic)"].fillna(0.0)
    for col in ["Monto antes de impuestos (MEMO)", "IEPS (MEMO)", "IVA (MEMO)", "Monto Total (MEMO)"]:
        if col in cruce.columns:
            cruce[col] = cruce[col].fillna(0.0)
    cruce["Amount Difference"] = (
        cruce["Gross claim amount (Panoptic)"] - cruce["Monto Total (MEMO)"]
    )
    cruce["Amount Status"] = np.where(
        cruce["Amount Difference"].abs() <= _AMOUNT_TOLERANCE, "OK", "Mismatch"
    )
    return cruce.drop(columns=["_merge"])


# ---------------------------------------------------------------------------
# Paso 4 — Duplicados Concepto-Año con clave extendida
# ---------------------------------------------------------------------------

def _paso4_duplicados(df_memo_ca: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """P4: Detecta duplicados en MEMO con clave de 7 campos.

    Clave: Num Proveedor + Año + Concepto + Num Categoria + Monto antes impuestos + IEPS + IVA.
    Devuelve (cruce_detalle_completo, solo_duplicados).
    """
    key = [c for c in _P4_KEY_COLS if c in df_memo_ca.columns]
    cruce_detalle = (
        df_memo_ca[key]
        .groupby(key, as_index=False)
        .size()
        .rename(columns={"size": "Veces en MEMO"})
    )
    if "Num Proveedor" in cruce_detalle.columns:
        cruce_detalle = cruce_detalle.rename(columns={"Num Proveedor": "Vendor number"})

    incon = cruce_detalle[cruce_detalle["Veces en MEMO"] > 1].copy()
    return cruce_detalle, incon


# ---------------------------------------------------------------------------
# Generación de plantilla de bulk update para Panoptic
# ---------------------------------------------------------------------------

def _generate_bulk_update_template(
    df_claims: pd.DataFrame,
    output_path: Path,
    template_path: Path | None = None,
) -> Path:
    """Genera el Excel de bulk update para Panoptic basado en la plantilla oficial.

    Escribe Project name, Claim number y Posting reference number en la Hoja 1
    a partir de la fila 2. La Hoja 2 se deja intacta tal como está en la plantilla.
    """
    from openpyxl import load_workbook

    resolved = template_path or _DEFAULT_TEMPLATE_PATH
    if not resolved.exists():
        raise FileNotFoundError(
            f"Plantilla no encontrada: {resolved}\n"
            f"Coloca 'Claim_Bulk_Update_Template.xlsx' en {resolved.parent}"
        )

    wb = load_workbook(resolved)
    ws = wb.worksheets[0]

    # Limpia filas de datos previas manteniendo encabezados (fila 1)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.value = None

    # Escribe los registros a partir de la fila 2
    cols = [c for c in _BULK_UPDATE_COLS if c in df_claims.columns]
    for i, row_data in enumerate(df_claims[cols].itertuples(index=False), start=2):
        for j, val in enumerate(row_data, start=1):
            ws.cell(row=i, column=j).value = val

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Lógica central de conciliación
# ---------------------------------------------------------------------------

def _reconcile_memo_from_df(
    df_pan_full: pd.DataFrame,
    memo_path: Path,
    memo_id: str,
    output_dir: Path,
    *,
    claims_accumulator: list[pd.DataFrame] | None = None,
    memo_vendor_sets: dict[str, set[str]] | None = None,
) -> Path:
    """Ejecuta los 4 pasos de conciliación para un memo.

    P1 — Asigna memo_id a claims sin Posting reference si el monto coincide con MEMO.
    P2 — Valida existencia de proveedores (df actualizado, filtrado por memo_id).
    P3 — Valida diferencias de monto (mismo df de P2).
    P4 — Detecta duplicados en MEMO con clave extendida de 7 campos.
    """
    df_pan = df_pan_full.copy()
    _ensure_tax_column(df_pan)

    df_memo_pf = _load_memo_prov_fase(memo_path)
    df_memo_ca = _load_memo_concepto_ano(memo_path)

    # P1 — Asignar memo a claims sin Posting reference (modifica df_pan in-place)
    df_p1_actualizados, df_p1_diferencias, df_claims = _paso1_asignar_memo(
        df_pan, df_memo_ca, memo_id, memo_vendor_sets or {}
    )

    # P1 segunda pasada — re-cruza diferencias incluyendo Stage="Invoice"
    df_p1_reejecucion = _paso1_reejecutar_diferencias(df_pan, df_memo_ca, df_p1_diferencias)

    if claims_accumulator is not None and not df_claims.empty:
        tagged = df_claims.copy()
        tagged.insert(0, "Memo", memo_id)
        claims_accumulator.append(tagged)

    # P2 + P3 — Cruce resumen con df ya actualizado por P1, filtrado por memo_id
    df_pan_memo = df_pan[df_pan["Posting reference number"] == memo_id].copy()
    if df_pan_memo.empty:
        raise ValueError(
            f"Sin registros en Panoptic para '{memo_id}' tras aplicar P1 — omitido."
        )

    cruce = _cruce_resumen(df_pan_memo, df_memo_pf)

    _prov_cols = [
        c for c in ["Vendor number", "Vendor name", "Vendor name (MEMO)", "Assigned to", "Vendor Status"]
        if c in cruce.columns
    ]
    incon_proveedores = cruce.loc[cruce["Vendor Status"] != "Match", _prov_cols].copy()

    # Asegura que Assigned to tenga valor aunque el vendor no tenga Posting reference asignado.
    # Busca en el DataFrame completo de Panoptic para capturar el auditor asignado.
    if "Assigned to" in df_pan.columns:
        if "Assigned to" not in incon_proveedores.columns:
            incon_proveedores["Assigned to"] = pd.NA
        assigned_lookup = df_pan.groupby("Vendor number")["Assigned to"].first()
        missing_mask = incon_proveedores["Assigned to"].isna()
        incon_proveedores.loc[missing_mask, "Assigned to"] = (
            incon_proveedores.loc[missing_mask, "Vendor number"].map(assigned_lookup)
        )
    incon_montos = cruce[
        (cruce["Vendor Status"] == "Match") & (cruce["Amount Status"] != "OK")
    ].copy()
    incon_montos = incon_montos[
        [c for c in _INCON_MONTOS_COLS if c in incon_montos.columns]
    ]

    # P4 — Duplicados en MEMO
    cruce_detalle, incon_concepto_ano = _paso4_duplicados(df_memo_ca)

    # Exportar — archivos por memo van en output_dir/memos/
    memos_dir = output_dir / "memos"
    memos_dir.mkdir(parents=True, exist_ok=True)
    output_path = memos_dir / f"conciliacion_{memo_id}.xlsx"

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        write_sheet(writer, df_p1_actualizados, "P1 Actualizados")
        write_sheet(writer, df_p1_diferencias,  "P1 Diferencias")
        write_sheet(writer, df_p1_reejecucion,  "P1 Reejecución")
        write_sheet(writer, incon_proveedores,  "Incons Proveedores")
        write_sheet(writer, incon_montos,       "Incons Montos")
        write_sheet(writer, incon_concepto_ano, "Incons Concepto-Año")
        write_sheet(writer, cruce,              "Cruce Resumen")
        write_sheet(writer, cruce_detalle,      "Cruce Detalle")
        style_workbook(writer.book)

    return output_path


# ---------------------------------------------------------------------------
# API pública
# ---------------------------------------------------------------------------

def process_panoptic_grouping(input_path: Path, output_dir: Path) -> Path:
    """Genera resumen agrupado por proveedor desde un XLSX de Panoptic."""
    df = pd.read_excel(input_path)
    _ensure_tax_column(df)

    summary = df.groupby(["Vendor number", "Vendor name"], as_index=False)[
        ["Net claim amount", "Total tax amount"]
    ].sum()
    summary["Gross claim amount"] = summary["Net claim amount"] + summary["Total tax amount"]
    summary = summary.rename(columns={
        "Net claim amount":  "Sum of Net claim amount",
        "Total tax amount":  "Sum of Total tax amount",
    })

    output_path = output_dir / f"resumen_proveedores_{input_path.name}"
    output_dir.mkdir(parents=True, exist_ok=True)
    summary.to_excel(output_path, index=False)
    return output_path


def reconcile_with_memo(
    raw_panoptic_path: Path,
    memo_path: Path,
    memo_id: str,
    output_dir: Path,
) -> Path:
    """Cruza un archivo Panoptic contra un MEMO individual."""
    df_pan = _load_panoptic(raw_panoptic_path)
    return _reconcile_memo_from_df(df_pan, memo_path, memo_id, output_dir)


@dataclass
class MemoResult:
    memo_id: str
    memo_path: Path
    output_path: Path | None = None
    error: str | None = None

    @property
    def ok(self) -> bool:
        return self.error is None


def _discover_memo_entries(memo_dir: Path) -> list[tuple[str, Path]]:
    """Devuelve lista de (memo_id, excel_path) ordenada por memo_id.

    Prioridad: subcarpetas Memo_* sobre archivos sueltos en la raíz.
    """
    entries: dict[str, Path] = {}

    for folder in memo_dir.iterdir():
        if not folder.is_dir():
            continue
        memo_id = _memo_id_from_name(folder.name)
        if not memo_id:
            continue
        xlsx = _find_memo_xlsx(folder)
        if xlsx:
            entries[memo_id] = xlsx

    for p in memo_dir.glob("Memo_*.xlsx"):
        if p.name.startswith("~$"):
            continue
        memo_id = _memo_id_from_path(p)
        if memo_id and memo_id not in entries:
            entries[memo_id] = p

    return sorted(entries.items(), key=lambda kv: kv[0])


def reconcile_all_memos(
    raw_panoptic_path: Path,
    memo_dir: Path,
    output_dir: Path,
    memo_from: int | None = None,
    memo_to: int | None = None,
    cancel_event=None,
) -> list[MemoResult]:
    """Carga Panoptic una sola vez y concilia contra todos los memos encontrados en memo_dir."""
    df_pan = _load_panoptic(raw_panoptic_path)

    entries = _discover_memo_entries(memo_dir)
    if not entries:
        raise FileNotFoundError(f"No se encontraron memos en {memo_dir}")

    if memo_from is not None or memo_to is not None:
        def _in_range(memo_id: str) -> bool:
            num_match = re.search(r"\d+", memo_id)
            if not num_match:
                return False
            num = int(num_match.group())
            if memo_from is not None and num < memo_from:
                return False
            if memo_to is not None and num > memo_to:
                return False
            return True
        entries = [(mid, p) for mid, p in entries if _in_range(mid)]

    # Pre-cargar vendor sets de todos los memos para la validación cruzada de P1.
    # Permite detectar si un claim con memo diferente asignado realmente pertenece a ese memo.
    memo_vendor_sets: dict[str, set[str]] = {}
    for mid, mpath in entries:
        try:
            df_pf = _load_memo_prov_fase(mpath)
            memo_vendor_sets[mid] = set(df_pf["Num Proveedor"].astype(str))
        except Exception:
            memo_vendor_sets[mid] = set()

    results: list[MemoResult] = []
    claims_accumulator: list[pd.DataFrame] = []

    for memo_id, memo_path in entries:
        if cancel_event and cancel_event.is_set():
            break
        result = MemoResult(memo_id=memo_id, memo_path=memo_path)
        try:
            result.output_path = _reconcile_memo_from_df(
                df_pan, memo_path, memo_id, output_dir,
                claims_accumulator=claims_accumulator,
                memo_vendor_sets=memo_vendor_sets,
            )
        except Exception as exc:
            result.error = str(exc)
        results.append(result)

    if claims_accumulator:
        all_claims = pd.concat(claims_accumulator, ignore_index=True)
        date_str = datetime.now().strftime("%Y%m%d")
        template_out = output_dir / f"Claim_Bulk_Update_{date_str}.xlsx"
        try:
            _generate_bulk_update_template(all_claims, template_out)
            print(f"[Bulk Update] Plantilla generada: {template_out}")
        except FileNotFoundError as exc:
            print(f"[Aviso] No se genero la plantilla de bulk update: {exc}")

    return results


def run_all(
    memo_dir: Path,
    output_dir: Path,
    memo_from: int | None = None,
    memo_to: int | None = None,
    download_dir: Path | None = None,
    panoptic_settings=None,
    view_name: str = "MONICA_3",
    print_fn=print,
) -> Path:
    """Descarga Panoptic, concilia todos los memos y genera el consolidado en un solo paso."""
    from ..panoptic.downloader import download_xlsx

    if download_dir is None:
        download_dir = output_dir / "panoptic"
    download_dir.mkdir(parents=True, exist_ok=True)

    print_fn("Paso 1/3 — Descargando Panoptic...")
    raw_path = download_xlsx(panoptic_settings, view_name=view_name)
    print_fn(f"  Panoptic descargado: {raw_path}")

    print_fn()
    print_fn("Paso 2/3 — Conciliando memos...")
    results = reconcile_all_memos(raw_path, memo_dir, output_dir, memo_from=memo_from, memo_to=memo_to)

    ok      = [r for r in results if r.ok]
    skipped = [r for r in results if not r.ok and "omitido" in (r.error or "")]
    failed  = [r for r in results if not r.ok and "omitido" not in (r.error or "")]

    for r in ok:
        print_fn(f"  OK  {r.memo_id}  ->  {r.output_path}")
    for r in skipped:
        print_fn(f"  --  {r.memo_id}  (sin registros en Panoptic, omitido)")
    for r in failed:
        print_fn(f"  ERR {r.memo_id}  {r.error}")
    print_fn(f"  Completados: {len(ok)} | Omitidos: {len(skipped)} | Errores: {len(failed)}")

    print_fn()
    print_fn("Paso 3/3 — Generando consolidado...")
    consolidated_path = consolidate_results(output_dir)
    print_fn(f"  Consolidado: {consolidated_path}")

    return consolidated_path


def consolidate_results(output_dir: Path, consolidated_path: Path | None = None) -> Path:
    """Lee todos los conciliacion_M*.xlsx de output_dir y genera un Excel consolidado.

    Hojas del archivo de salida:
    - Resumen          : una fila por memo con conteos y estado general.
    - P1 Actualizados  : proveedores asignados al memo vía P1 (todos los memos).
    - P1 Diferencias   : proveedores con diferencia en P1 (todos los memos).
    - Incons Proveedores, Incons Montos, Incons Concepto-Año: inconsistencias consolidadas.
    """
    memos_dir = output_dir / "memos"
    archivos = sorted(memos_dir.glob("conciliacion_M*.xlsx")) if memos_dir.exists() \
               else sorted(output_dir.glob("conciliacion_M*.xlsx"))
    if not archivos:
        raise FileNotFoundError(
            f"No se encontraron archivos conciliacion_M*.xlsx en {memos_dir} ni en {output_dir}"
        )

    resumen_rows:        list[dict] = []
    all_p1_actualizados: list[pd.DataFrame] = []
    all_p1_diferencias:  list[pd.DataFrame] = []
    all_p1_reejecucion:  list[pd.DataFrame] = []
    all_incons_prov:     list[pd.DataFrame] = []
    all_incons_montos:   list[pd.DataFrame] = []
    all_incons_conc:     list[pd.DataFrame] = []

    for path in archivos:
        memo_match = re.search(r"conciliacion_(M\d+)", path.stem)
        if not memo_match:
            continue
        mid = memo_match.group(1)

        try:
            xl     = pd.ExcelFile(path)
            sheets = xl.sheet_names

            def _read(name: str) -> pd.DataFrame:
                if name in sheets:
                    return pd.read_excel(xl, sheet_name=name, header=HEADER_ROWS)
                return pd.DataFrame()

            d_p1_act      = _read("P1 Actualizados")
            d_p1_dif      = _read("P1 Diferencias")
            d_p1_reej     = _read("P1 Reejecución")
            d_proveedores = _read("Incons Proveedores")
            d_montos      = _read("Incons Montos")
            d_conc        = _read("Incons Concepto-Año")
            d_cruce       = _read("Cruce Resumen")
        except Exception as exc:
            resumen_rows.append({"Memo": mid, "Estado": f"Error al leer: {exc}"})
            continue

        pan_provs  = (
            d_cruce[d_cruce["Vendor Status"] != "Missing in Panoptic"]["Vendor number"].nunique()
            if not d_cruce.empty and "Vendor Status" in d_cruce.columns else 0
        )
        memo_provs = (
            d_cruce[d_cruce["Vendor Status"] != "Missing in MEMO"]["Vendor number"].nunique()
            if not d_cruce.empty and "Vendor Status" in d_cruce.columns else 0
        )

        cnt_p1_act  = len(d_p1_act)
        cnt_p1_dif  = len(d_p1_dif)
        cnt_p1_reej = len(d_p1_reej)
        cnt_prov    = len(d_proveedores)
        cnt_montos  = len(d_montos)
        cnt_conc    = len(d_conc)
        estado = (
            "OK - Sin inconsistencias"
            if (cnt_prov == 0 and cnt_montos == 0 and cnt_conc == 0)
            else "Con inconsistencias"
        )

        resumen_rows.append({
            "Memo":                  mid,
            "P1 Actualizados":       cnt_p1_act,
            "P1 Diferencias":        cnt_p1_dif,
            "P1 Reejec Invoice":     cnt_p1_reej,
            "Proveedores Panoptic":  pan_provs,
            "Proveedores MEMO":      memo_provs,
            "Proveedores coinciden": "Si" if pan_provs == memo_provs else "No",
            "Incons Proveedores":    cnt_prov,
            "Incons Montos":         cnt_montos,
            "Incons Concepto/Año":   cnt_conc,
            "Estado":                estado,
        })

        def _tag(df: pd.DataFrame) -> pd.DataFrame:
            d = df.copy()
            d.insert(0, "Memo", mid)
            return d

        if cnt_p1_act  > 0: all_p1_actualizados.append(_tag(d_p1_act))
        if cnt_p1_dif  > 0: all_p1_diferencias.append(_tag(d_p1_dif))
        if cnt_p1_reej > 0: all_p1_reejecucion.append(_tag(d_p1_reej))
        if cnt_prov    > 0: all_incons_prov.append(_tag(d_proveedores))
        if cnt_montos  > 0: all_incons_montos.append(_tag(d_montos))
        if cnt_conc    > 0: all_incons_conc.append(_tag(d_conc))

    def _concat(frames: list[pd.DataFrame]) -> pd.DataFrame:
        return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()

    if consolidated_path is None:
        consolidated_path = output_dir / "Resumen_conciliacion_memo_panoptic.xlsx"

    with pd.ExcelWriter(consolidated_path, engine="openpyxl") as writer:
        write_sheet(writer, pd.DataFrame(resumen_rows),   "Resumen")
        write_sheet(writer, _concat(all_p1_actualizados), "P1 Actualizados")
        write_sheet(writer, _concat(all_p1_diferencias),  "P1 Diferencias")
        write_sheet(writer, _concat(all_p1_reejecucion),  "P1 Reejecución")
        write_sheet(writer, _concat(all_incons_prov),     "Incons Proveedores")
        write_sheet(writer, _concat(all_incons_montos),   "Incons Montos")
        write_sheet(writer, _concat(all_incons_conc),     "Incons Concepto-Año")
        style_workbook(writer.book)

    return consolidated_path
